from io import BytesIO

from app.core.auth import issue_access_token
from app.core.config import Settings
from docx import Document as DocxDocument
from fastapi.testclient import TestClient


def _headers(*permissions: str) -> dict[str, str]:
    token = issue_access_token("doc-user", ["auditor"], list(permissions), ["*"], Settings().auth_secret, 900)
    return {"Authorization": f"Bearer {token}"}


def test_markdown_upload_parses_parts_and_preserves_source_metadata(client: TestClient) -> None:
    response = client.post(
        "/api/rag/documents",
        files={"file": ("rules.md", b"# Street lights\n\nReplace failed lamps within 24 hours.", "text/markdown")},
        data={
            "facility_type": "street_light",
            "region": "DISTRICT-01",
            "effective_date": "2026-01-01",
            "expiry_date": "2027-01-01",
            "version": "v2",
            "source_type": "internal",
        },
        headers=_headers("document:upload"),
    )
    assert response.status_code == 201
    body = response.json()
    assert body["item"]["status"] == "PARSED"
    assert body["item"]["original_filename"] == "rules.md"
    assert body["item"]["version"] == "v2"
    assert body["parts"][0]["heading"] == "Street lights"
    assert "24 hours" in body["parts"][0]["text"]


def test_document_duplicate_and_unsafe_filename_are_rejected(client: TestClient) -> None:
    headers = _headers("document:upload")
    payload = {"file": ("rules.txt", b"same content", "text/plain")}
    first = client.post("/api/rag/documents", files=payload, headers=headers)
    assert first.status_code == 201
    duplicate = client.post("/api/rag/documents", files=payload, headers=headers)
    assert duplicate.status_code == 409
    assert duplicate.json()["code"] == "DUPLICATE_DOCUMENT"
    unsafe = client.post(
        "/api/rag/documents",
        files={"file": ("../secrets.txt", b"secret", "text/plain")},
        headers=headers,
    )
    assert unsafe.status_code == 400
    assert unsafe.json()["code"] == "UNSAFE_FILENAME"


def test_html_upload_extracts_table_text_and_metadata_endpoint_is_readable(client: TestClient) -> None:
    html = b"<html><head><title>Rules</title></head><body><h1>Inspection</h1><table><tr><th>Level</th><th>SLA</th></tr><tr><td>P1</td><td>4h</td></tr></table></body></html>"
    upload = client.post(
        "/api/rag/documents",
        files={"file": ("rules.html", html, "text/html")},
        headers=_headers("document:upload"),
    )
    assert upload.status_code == 201
    document_id = upload.json()["item"]["id"]
    assert any("Level" in part["text"] and "4h" in part["text"] for part in upload.json()["parts"])
    detail = client.get(f"/api/rag/documents/{document_id}", headers=_headers("document:read"))
    assert detail.status_code == 200
    assert detail.json()["item"]["source_type"] == "internal"


def test_corrupt_pdf_is_recorded_as_failed_parse(client: TestClient) -> None:
    response = client.post(
        "/api/rag/documents",
        files={"file": ("broken.pdf", b"%PDF-not-a-real-file", "application/pdf")},
        headers=_headers("document:upload"),
    )
    assert response.status_code == 422
    assert response.json()["code"] == "DOCUMENT_PARSE_FAILED"
    document_id = response.json()["document_id"]
    detail = client.get(f"/api/rag/documents/{document_id}", headers=_headers("document:read"))
    assert detail.status_code == 200
    assert detail.json()["item"]["status"] == "FAILED"
    assert detail.json()["item"]["error_reason"]


def test_docx_upload_extracts_heading_and_table(client: TestClient) -> None:
    document = DocxDocument()
    document.add_heading("Maintenance", level=1)
    document.add_paragraph("Inspect the cabinet before repair.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Priority"
    table.cell(0, 1).text = "SLA"
    table.cell(1, 0).text = "P1"
    table.cell(1, 1).text = "4h"
    content = BytesIO()
    document.save(content)
    response = client.post(
        "/api/rag/documents",
        files={
            "file": (
                "maintenance.docx",
                content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=_headers("document:upload"),
    )
    assert response.status_code == 201
    assert response.json()["parts"][0]["heading"] == "Maintenance"
    assert any("Priority" in part["text"] and "4h" in part["text"] for part in response.json()["parts"])
